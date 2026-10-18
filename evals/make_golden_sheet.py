"""Build evals/golden_50_labeling.xlsx for the human labeler (no model calls).

Star rating is deliberately hidden so labels come from the text, as the contract requires.
After labeling, run: python3 evals/import_golden.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from check_submission import INTENTS, TOPICS, csv_rows  # noqa: E402

COLS = ["#", "review_id", "review_text", "topic", "intent", "severity", "sentiment", "evidence_quote",
        "entities", "needs_review", "alt_topic", "alt_intent", "alt_severity", "notes"]
WIDTHS = [4, 12, 70, 13, 13, 9, 10, 40, 18, 12, 12, 12, 11, 30]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "label_here"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, row in enumerate(csv_rows(ROOT / "data/golden_50_to_label.csv"), 1):
        ws.append([i, row["review_id"], row["review_text"]] + [None] * (len(COLS) - 3))
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    for r in ws.iter_rows(min_row=2):
        r[2].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    n = ws.max_row

    def dv(options, cols):
        v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(v)
        for col in cols:
            v.add(f"{col}2:{col}{n}")

    dv(TOPICS, ["D", "K"])
    dv(INTENTS, ["E", "L"])
    dv(["1", "2", "3", "4", "5"], ["F", "M"])
    dv(["-1", "-0.5", "0", "0.5", "1"], ["G"])
    dv(["FALSE", "TRUE"], ["J"])

    h = wb.create_sheet("how_to")
    lines = [
        "HOW TO LABEL (about 40 minutes)",
        "Required for every row: topic, intent, severity, sentiment (dropdowns).",
        "Optional: evidence_quote (copy an exact phrase from the review), entities (comma-separated features named), needs_review (TRUE if genuinely ambiguous).",
        "alt_topic / alt_intent / alt_severity: fill ONLY when a second answer is equally defensible — the grader accepts multiple labels for ambiguous cases.",
        "Judge the text only. Star ratings are hidden on purpose.",
        "Full definitions + examples: prompts/labels_v1.md. Quick reference below.",
        "",
        "TOPIC: access=login/account | usability=navigation, controls, shuffle/queue, ADS | playback=won't play, crashes, lag, audio, 'no internet'",
        "       downloads=offline/saved music | catalog=missing songs, search, recommendations, lyrics availability",
        "       billing=price, charges, subscription, paywalls, PREMIUM-ONLY features | support=contacting support | other=general praise/criticism, unrelated",
        "  Highest-severity problem wins; tie -> first mentioned. Positive review -> first praised feature; general praise -> other.",
        "INTENT (first that applies): cancellation (leaving/uninstalling/threatening) > complaint (any negative, incl. mixed) > request > praise > unclear",
        "  Bare boycott slogans / meaningless text -> unclear. 'Bad app' -> complaint.",
        "SEVERITY: 1 no problem/praise/request | 2 annoyance, generic criticism | 3 degraded/restricted, workaround remains",
        "          4 core task blocked (can't log in / can't play) | 5 explicit financial, privacy or data harm (angry tone or price alone is NOT 5)",
        "SENTIMENT: -1 very negative, -0.5 negative, 0 neutral/mixed, 0.5 positive, 1 very positive",
        "",
        "Save the file in place when done (keep the name golden_50_labeling.xlsx).",
    ]
    for line in lines:
        h.append([line])
    h.column_dimensions["A"].width = 160
    wb.move_sheet("how_to", offset=-1)
    out = ROOT / "evals/golden_50_labeling.xlsx"
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
