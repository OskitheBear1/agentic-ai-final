"""Convert the human-labeled spreadsheet into evaluation files (no model calls).

Outputs:
  evals/golden_50_labeled.csv   human labels in the course's golden CSV format
  evals/gold_benchmark.json     same labels in check_submission.py --gold format (accepted_* lists)

These files are never read by any model-calling stage (enforced by a test in P2).
"""

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from check_submission import INTENTS, TOPICS, csv_rows  # noqa: E402

GOLD_FIELDS = ["review_id", "review_text", "review_rating", "review_likes", "app_version", "review_timestamp",
               "topic", "intent", "sentiment", "severity", "entities", "evidence_quote", "needs_review"]


def cell(v):
    return "" if v is None else str(v).strip()


def main():
    source = {r["review_id"]: r for r in csv_rows(ROOT / "data/golden_50_to_label.csv")}
    src_xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "evals/golden_50_labeled.xlsx"
    wb = load_workbook(src_xlsx)
    ws = wb["label_here"] if "label_here" in wb.sheetnames else wb.worksheets[0]
    header = [c.value for c in ws[1]]
    text_col = header.index("review_text")
    all_rows = list(ws.iter_rows(min_row=1, values_only=True))

    def resolve(v):
        # A formula like "=B3" means "the whole review text of spreadsheet row 3".
        v = cell(v)
        if v.startswith("=") and v[1:2].isalpha() and v[2:].isdigit():
            return cell(all_rows[int(v[2:]) - 1][ord(v[1].upper()) - 65])
        return v
    errors, out_rows, cases = [], [], []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values))
        rid = cell(row["review_id"])
        if not rid:
            continue
        src = source[rid]
        topic, intent = cell(row["topic"]), cell(row["intent"])
        try:
            severity = int(float(cell(row["severity"])))
            sentiment = float(cell(row["sentiment"]))
        except ValueError:
            errors.append(f"{rid}: severity/sentiment missing")
            continue
        if topic not in TOPICS or intent not in INTENTS or not 1 <= severity <= 5 or not -1 <= sentiment <= 1:
            errors.append(f"{rid}: invalid label {topic}/{intent}/{severity}/{sentiment}")
            continue
        quote = resolve(row["evidence_quote"])
        if quote and quote not in src["review_text"]:
            errors.append(f"{rid}: evidence_quote is not an exact substring")
        entities = [e.strip() for e in cell(row["entities"]).split(",") if e.strip()]
        needs = cell(row["needs_review"]).upper() in {"TRUE", "1", "YES"}
        out_rows.append({**{k: src[k] for k in GOLD_FIELDS[:6]}, "topic": topic, "intent": intent, "sentiment": sentiment,
                         "severity": severity, "entities": json.dumps(entities), "evidence_quote": quote,
                         "needs_review": str(needs).lower()})
        alt = lambda name, cast: [cast(cell(row.get(name)))] if cell(row.get(name)) else []  # noqa: E731
        cases.append({"review_id": rid, "status": "approved", "reviewer": "student",
                      "accepted_topic": [topic] + [t for t in alt("alt_topic", str) if t != topic],
                      "accepted_intent": [intent] + [t for t in alt("alt_intent", str) if t != intent],
                      "accepted_severity": [severity] + [s for s in alt("alt_severity", lambda x: int(float(x))) if s != severity],
                      "sentiment": sentiment, "needs_review": needs})
    if errors:
        print("Fix these rows in the spreadsheet, then rerun:\n  " + "\n  ".join(errors))
        sys.exit(1)
    if len(cases) != 50:
        print(f"Only {len(cases)} of 50 rows labeled.")
        sys.exit(1)
    with (ROOT / "evals/golden_50_labeled.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=GOLD_FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(out_rows)
    (ROOT / "evals/gold_benchmark.json").write_text(json.dumps({"version": "student-golden-50-v1", "cases": cases}, indent=2) + "\n")
    ambiguous = sum(len(c["accepted_topic"]) + len(c["accepted_intent"]) + len(c["accepted_severity"]) > 3 for c in cases)
    print(f"Imported 50 labels ({ambiguous} marked ambiguous with alternates).")


if __name__ == "__main__":
    main()
