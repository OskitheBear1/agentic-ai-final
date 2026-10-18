"""Apply the human label audit (evals/golden_audit_v2.xlsx) as a NEW golden version; v1 files stay untouched.

    .venv/bin/python evals/apply_golden_audit.py

Selection of audited cases was rule-based, made by code with no model predictions shown: intent praise/unclear with
severity != 1, and intent request with severity != 1. Every change was decided by the student. Writes:
  evals/golden_50_labeled_v2.csv, evals/gold_benchmark_v2.json (version student-golden-50-v1.2), evals/label_audit.md
"""
import csv
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
FIELDS = {"topic": 7, "intent": 8, "severity": 9}  # NEW_* column index in the audit sheet


def main():
    ws = load_workbook(ROOT / "evals/golden_audit_v2.xlsx").active
    rows = list(csv.DictReader(open(ROOT / "evals/golden_50_labeled.csv", encoding="utf-8")))
    by_id = {r["review_id"]: r for r in rows}
    bench = json.loads((ROOT / "evals/gold_benchmark.json").read_text())
    cases = {c["review_id"]: c for c in bench["cases"]}
    changes = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rid, why, note = r[1], r[6], r[10]
        orig = by_id[rid]  # "before" comes from the v1 CSV, never from the sheet's display columns
        for field, col in FIELDS.items():
            new = r[col]
            if new in (None, ""):
                continue
            new = str(int(new)) if field == "severity" else str(new).strip()
            if new == orig[field]:
                continue
            changes.append({"review_id": rid, "text": orig["review_text"], "field": field, "old": orig[field], "new": new,
                            "rule": why, "note": note or ""})
            orig[field] = new
            c = cases[rid]
            key = f"accepted_{field}"
            val = int(new) if field == "severity" else new
            c[key] = [val] + [v for v in c[key] if v != val and str(v) != changes[-1]["old"]]
    with (ROOT / "evals/golden_50_labeled_v2.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    bench["version"] = "student-golden-50-v1.2"
    bench.setdefault("corrections", []).append({
        "date": "2026-10-18", "change": f"{len(changes)} label changes from a rule-consistency audit; see evals/label_audit.md",
        "cases": len({c['review_id'] for c in changes}), "decided_by": "student",
        "prompted_by": "external review flagged rule violations; cases selected by a code rule check, model predictions not shown"})
    (ROOT / "evals/gold_benchmark_v2.json").write_text(json.dumps(bench, indent=2) + "\n")
    lines = ["# Golden-set label audit (v1 → v2)", "",
             "**Why:** an external review found human labels that break the shared rules (praise/unclear rated severity 2; "
             "requests that report a failure). **Selection:** a code rule check over all 50 labels (intent praise/unclear "
             "or request with severity ≠ 1) found 15 cases. Model predictions were **not** shown during the audit. "
             "**Decisions:** the student, using `prompts/labels_v1.md`. The audit was prompted by reviewing disagreements, so "
             "v2 is not independent of the pipeline's results; both versions are reported.", "",
             f"**{len(changes)} label changes across {len({c['review_id'] for c in changes})} reviews.** The v1 files are unchanged.", "",
             "| review (start) | field | v1 → v2 | rule |", "|---|---|---|---|"]
    for c in changes:
        lines.append(f"| {c['text'][:60].replace('|', '/').replace(chr(10), ' ')} | {c['field']} | {c['old']} → {c['new']} | "
                     f"{c['rule']}{' ' + c['note'] if c['note'] else ''} |")
    (ROOT / "evals/label_audit.md").write_text("\n".join(lines) + "\n")
    print(f"{len(changes)} changes -> evals/golden_50_labeled_v2.csv, evals/gold_benchmark_v2.json, evals/label_audit.md")


if __name__ == "__main__":
    main()
